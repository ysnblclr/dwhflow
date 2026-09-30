"""Yeni biçim okuyucuları: eski Office, sunum, ODF, RTF, HTML, e-posta, arşiv, diyagramlar,
kurum şablonları ve genel dönüştürücü.

Her biçim için: ayrıştırma, içerik tespiti (ve ilgisiz dosyaları çalmaması), bozuk girdide
istisnasız ``error``, önizleme konumlarının okuyucu konumlarıyla eşleşmesi.
"""

from __future__ import annotations

import io
import json
import sys
import time
import types
import zipfile
from pathlib import Path

import pytest
import yaml

from dwhflow.ingest import ACCEPTED, PLUGIN_ERRORS, catalog, ingest_files, preview_document
from dwhflow.ingest.config import DEFAULT, IngestConfig, using
from dwhflow.ingest.formats import _acma, _diyagram, eski_ofis, mermaid, opendocument, sablon
from dwhflow.ingest.formats._html import html_items
from dwhflow.ingest.formats._ortak import FormatError
from dwhflow.ingest.formats._rtf import rtf_items
from dwhflow.ingest.formats.graphviz_dot import dot_pages
from dwhflow.ingest.registry import REGISTRY, register_expander, resolve

from .bicim_yardimci import (
    FORMATS,
    TEMPLATES,
    assert_locators_in_preview,
    docx_bytes,
    generator,
    source_locators,
    summary,
)

NEW_FORMATS = {"xls", "doc", "pptx", "odt", "ods", "rtf", "html", "eposta", "msg", "arsiv", "visio", "drawio",
               "bpmn", "mermaid", "dot", "sablon", "genel"}


@pytest.fixture
def scratch_registry():
    """Testin kaydettiği biçimleri geri alır; genel kayıt kirlenmesin."""
    formats, by_suffix = dict(REGISTRY.formats), dict(REGISTRY.by_suffix)
    expanders, unsupported = dict(REGISTRY.expanders), dict(REGISTRY.unsupported)
    yield REGISTRY
    REGISTRY.formats.clear(), REGISTRY.formats.update(formats)
    REGISTRY.by_suffix.clear(), REGISTRY.by_suffix.update(by_suffix)
    REGISTRY.expanders.clear(), REGISTRY.expanders.update(expanders)
    REGISTRY.unsupported.clear(), REGISTRY.unsupported.update(unsupported)


@pytest.fixture
def no_soffice(monkeypatch):
    """LibreOffice yokmuş gibi (test makinesinde kurulu olsa bile belirlenimci)."""
    monkeypatch.setattr(eski_ofis, "find_soffice", lambda: None)


def only(result, file: str):
    return next(r for r in result.reports if r.file == file)


# --------------------------------------------------------------------------- #
# Kayıt
# --------------------------------------------------------------------------- #


class TestRegistration:
    def test_plugins_load_without_errors(self):
        assert not [e for e in PLUGIN_ERRORS if not e.startswith(("airflow", "controlm", "datastage"))]

    def test_catalog_lists_new_formats(self):
        keys = {item["key"] for item in catalog()}
        assert NEW_FORMATS <= keys

    @pytest.mark.parametrize("suffix,key", [
        (".xls", "xls"), (".doc", "doc"), (".pptx", "pptx"), (".odt", "odt"), (".ods", "ods"), (".rtf", "rtf"),
        (".vsdx", "visio"), (".drawio", "drawio"), (".bpmn", "bpmn"), (".mmd", "mermaid"), (".dot", "dot"),
        (".gv", "dot"), (".eml", "eposta"), (".msg", "msg"), (".zip", "arsiv")])
    def test_suffix(self, tmp_path, suffix, key):
        path = tmp_path / f"x{suffix}"
        path.write_bytes(b"x")
        spec, how, _ = resolve(path, DEFAULT)
        assert spec.key == key and how == "uzanti"

    def test_accepted_for_upload(self):
        for suffix in (".xls", ".doc", ".pptx", ".odt", ".ods", ".rtf", ".vsdx", ".drawio", ".bpmn", ".mmd",
                       ".dot", ".eml", ".msg", ".zip"):
            assert suffix in ACCEPTED

    def test_html_is_read_but_not_directly_uploadable(self, tmp_path):
        """Web arayüzü girdileri geri sunduğu için HTML yüklemesi kapalı (depolanmış XSS); okuma açık."""
        assert not {".html", ".htm", ".xhtml", ".mht", ".mhtml"} & set(ACCEPTED)
        path = tmp_path / "sayfa.HTML"
        path.write_text("<p>x</p>", encoding="utf-8")
        assert resolve(path, DEFAULT)[0].key == "html"

    def test_other_archives_explained(self, tmp_path):
        path = tmp_path / "paket.rar"
        path.write_bytes(b"Rar!\x1a\x07\x00")
        report = ingest_files([path]).reports[0]
        assert report.status == "unsupported" and "zip" in report.notes[0]


# --------------------------------------------------------------------------- #
# Eski Office: .xls
# --------------------------------------------------------------------------- #


class TestXls:
    def test_values_and_sheets(self):
        sheets = dict(eski_ofis.xls_sheets(FORMATS / "kart_sttm_eski.xls"))
        assert list(sheets) == ["STTM", "Job Bağımlılıkları", "Notlar"]
        header = sheets["STTM"][2]
        assert header == ["Kaynak Tablo", "Hedef Tablo", "Job", "Yükleme Tipi", "Sürüm"]
        assert sheets["STTM"][3][4] == 1 and isinstance(sheets["STTM"][3][4], int)  # 1.0 → 1
        assert sheets["STTM"][6][2] is None  # boş hücre

    def test_same_semantics_as_excel(self, tmp_path):
        """Aynı satırlar .xlsx ve .xls olarak aynı modeli vermeli."""
        import openpyxl

        rows = [["Kaynak Tablo", "Hedef Tablo", "Job"], ["STG_A", "ODS_A", "SP_A"], ["ODS_A", "F_A", None]]
        xls = tmp_path / "a.xls"
        xls.write_bytes(generator().xls_bytes([("STTM", rows)]))
        book = openpyxl.Workbook()
        book.active.title = "STTM"
        for row in rows:
            book.active.append(row)
        book.save(tmp_path / "a.xlsx")
        old, new = ingest_files([xls]), ingest_files([tmp_path / "a.xlsx"])
        assert summary(old) == summary(new)
        assert source_locators(old, "a.xls") == {"STTM · satır 1", "STTM · satır 2"}

    def test_preview_locators(self):
        path = FORMATS / "kart_sttm_eski.xls"
        assert_locators_in_preview(path, ingest_files([path]))

    @pytest.mark.parametrize("content,note", [
        (b"<html><body><table><tr><td>Kaynak Tablo</td><td>Hedef Tablo</td></tr>"
         b"<tr><td>STG_A</td><td>ODS_A</td></tr></table></body></html>", "HTML"),
        (b"Kaynak Tablo\tHedef Tablo\nSTG_A\tODS_A\n", "CSV"),
        (b'<?xml version="1.0"?><Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" '
         b'xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet ss:Name="S"><Table>'
         b'<Row><Cell><Data ss:Type="String">Kaynak Tablo</Data></Cell><Cell><Data>Hedef Tablo</Data></Cell></Row>'
         b'<Row ss:Index="3"><Cell><Data>STG_A</Data></Cell><Cell ss:Index="2"><Data>ODS_A</Data></Cell></Row>'
         b"</Table></Worksheet></Workbook>", "XML Elektronik Tablo"),
    ])
    def test_disguised_xls(self, tmp_path, content, note):
        path = tmp_path / "rapor.xls"
        path.write_bytes(content)
        result = ingest_files([path])
        assert {("STG_A", "ODS_A")} == summary(result)["links"]
        assert any(note in n for n in result.reports[0].notes)

    def test_xlsx_saved_as_xls(self, tmp_path):
        import openpyxl

        book = openpyxl.Workbook()
        book.active.append(["Kaynak Tablo", "Hedef Tablo"])
        book.active.append(["STG_A", "ODS_A"])
        buffer = io.BytesIO()
        book.save(buffer)
        path = tmp_path / "yanlis.xls"
        path.write_bytes(buffer.getvalue())
        assert summary(ingest_files([path]))["links"] == {("STG_A", "ODS_A")}

    @pytest.mark.parametrize("content", [
        generator().CFB_SIGNATURE + b"\x00" * 600, b"\x00\x01\x02\x03" * 50,
        generator().cfb([("Baska", b"veri")])])
    def test_corrupt_xls_is_error(self, tmp_path, content):
        path = tmp_path / "bozuk.xls"
        path.write_bytes(content)
        report = ingest_files([path]).reports[0]
        assert report.status == "error" and report.notes


# --------------------------------------------------------------------------- #
# Eski Office: .doc
# --------------------------------------------------------------------------- #

FLOW = "STG_A tablosundan ODS_A tablosuna SP_A prosedürüyle aktarılır."


class TestDoc:
    def test_fallback_text_extraction(self, tmp_path, no_soffice):
        path = tmp_path / "eski.doc"
        path.write_bytes(generator().doc_bytes(["1 Yükleme", FLOW, "Başka bir paragraf."]))
        result = ingest_files([path])
        report = result.reports[0]
        assert report.status == "ok" and "LibreOffice bulunamadı" in report.notes[0] and ".docx" in report.notes[0]
        assert summary(result)["links"] == {("STG_A", "ODS_A")}
        assert source_locators(result, "eski.doc") == {"§ 1 Yükleme · paragraf 1"}
        assert_locators_in_preview(path, result)

    def test_word_special_characters(self):
        items = eski_ofis.word_text_items(
            'Bkz. \x13 HYPERLINK "http://x" \x14wiki sayfası\x15 ve\x0bsonrası\r'
            "Kaynak Tablo\x07Hedef Tablo\x07\x07STG_A\x07ODS_A\x07\x07\r")
        assert [i.text for i in items] == ["Bkz. wiki sayfası ve\nsonrası", "Kaynak Tablo | Hedef Tablo",
                                           "STG_A | ODS_A"]

    def test_printable_runs_when_fib_is_broken(self, tmp_path, no_soffice):
        broken = generator().cfb([("WordDocument", b"\x00" * 64 + FLOW.encode("utf-16-le"))])
        path = tmp_path / "bozuk_fib.doc"
        path.write_bytes(broken)
        result = ingest_files([path])
        assert summary(result)["links"] == {("STG_A", "ODS_A")}

    def test_encrypted_doc_is_error(self, tmp_path, no_soffice):
        data = bytearray(generator().doc_bytes([FLOW]))
        stream = data.find(b"\xec\xa5\xc1\x00")  # FIB imzası
        data[stream + 0x0A] |= 0x01  # fEncrypted (0x0100) bayrağının yüksek baytı
        data[stream + 0x0B] |= 0x01
        path = tmp_path / "sifreli.doc"
        path.write_bytes(bytes(data))
        report = ingest_files([path]).reports[0]
        assert report.status == "error" and "şifreli" in report.notes[0]

    @pytest.mark.parametrize("content,note", [
        (b"{\\rtf1\\ansi " + FLOW.replace("ı", "\\'fd").replace("ü", "\\'fc").encode("ascii") + b"\\par}", "RTF"),
        (b"MIME-Version: 1.0\r\nContent-Type: multipart/related; boundary=\"b\"\r\n\r\n--b\r\n"
         b"Content-Type: text/html; charset=utf-8\r\n\r\n<p>" + FLOW.encode() + b"</p>\r\n--b--\r\n", "MHTML"),
    ])
    def test_other_contents_named_doc(self, tmp_path, content, note):
        path = tmp_path / "disa_aktarim.doc"
        path.write_bytes(content)
        result = ingest_files([path])
        assert summary(result)["links"] == {("STG_A", "ODS_A")}
        assert any(note in n for n in result.reports[0].notes)

    def test_docx_named_doc(self, tmp_path):
        path = tmp_path / "aslinda.doc"
        path.write_bytes(docx_bytes([(None, FLOW)]))
        result = ingest_files([path])
        assert summary(result)["links"] == {("STG_A", "ODS_A")} and ".docx" in result.reports[0].notes[0]

    def test_not_a_word_file(self, tmp_path):
        path = tmp_path / "rastgele.doc"
        path.write_bytes(b"\x00\x01\x02 bu bir word dosyasi degil")
        report = ingest_files([path]).reports[0]
        assert report.status == "error" and "tanınmadı" in report.notes[0]

    def test_libreoffice_conversion_used(self, tmp_path, monkeypatch):
        commands = []

        def fake_run(command, timeout):
            commands.append((command, timeout))
            outdir = Path(command[command.index("--outdir") + 1])
            (outdir / "girdi.docx").write_bytes(docx_bytes([("Heading1", "Yükleme"), (None, FLOW)]))
            return 0

        monkeypatch.setattr(eski_ofis, "find_soffice", lambda: "soffice")
        monkeypatch.setattr(eski_ofis, "run_command", fake_run)
        path = tmp_path / "Kart Tasarım.doc"
        path.write_bytes(generator().doc_bytes(["çöp"]))
        result = ingest_files([path])
        command, timeout = commands[0]
        assert command[0] == "soffice" and "--headless" in command and timeout == eski_ofis.CONVERT_TIMEOUT
        assert any(part.startswith("-env:UserInstallation=file:") for part in command)  # ayrık profil
        assert "LibreOffice ile" in result.reports[0].notes[0]
        # Word okuyucusunun konumları ve özgün dosya adı
        assert source_locators(result, "Kart Tasarım.doc") == {"§ Yükleme · paragraf 1"}
        assert summary(result)["links"] == {("STG_A", "ODS_A")}
        assert_locators_in_preview(path, result)

    @pytest.mark.parametrize("behaviour", ["no_output", "timeout", "empty_document"])
    def test_libreoffice_failure_falls_back(self, tmp_path, monkeypatch, behaviour):
        def fake_run(command, timeout):
            if behaviour == "timeout":
                raise FormatError("LibreOffice 120 sn içinde bitmedi")
            if behaviour == "empty_document":
                outdir = Path(command[command.index("--outdir") + 1])
                (outdir / "girdi.docx").write_bytes(docx_bytes([]))
            return 1

        monkeypatch.setattr(eski_ofis, "find_soffice", lambda: "soffice")
        monkeypatch.setattr(eski_ofis, "run_command", fake_run)
        path = tmp_path / "eski.doc"
        path.write_bytes(generator().doc_bytes([FLOW]))
        result = ingest_files([path])
        assert result.reports[0].status == "ok"
        assert "dönüştürülemedi" in result.reports[0].notes[0]
        assert summary(result)["links"] == {("STG_A", "ODS_A")}

    def test_run_command_kills_on_timeout(self):
        started = time.monotonic()
        with pytest.raises(FormatError, match="bitmedi"):
            eski_ofis.run_command([sys.executable, "-c", "import time; time.sleep(60)"], 1)
        assert time.monotonic() - started < 30

    def test_soffice_can_be_disabled(self, monkeypatch):
        monkeypatch.setenv("DWHFLOW_LIBREOFFICE", "yok")
        assert eski_ofis.find_soffice() is None


# --------------------------------------------------------------------------- #
# PowerPoint, OpenDocument
# --------------------------------------------------------------------------- #


class TestPptx:
    def test_slide_order_follows_presentation(self):
        _, blocks = preview_document(FORMATS / "kart_mimari_sunum.pptx")
        headings = [b["text"] for b in blocks if b["kind"] == "heading"]
        assert headings == ["Slayt 1 — Kart İşlemleri Veri Ambarı Mimarisi", "Slayt 2 — Mart katmanı"]
        assert not any(b["text"] == "2" for b in blocks if b["kind"] == "paragraph")  # slayt no yer tutucusu

    def test_locators(self):
        path = FORMATS / "kart_mimari_sunum.pptx"
        result = ingest_files([path])
        used = source_locators(result, path.name)
        assert {"slayt 2 · paragraf 2", "slayt 2 · not 1"} <= used
        assert any(u.startswith("slayt 2 · tablo 1 · satır ") for u in used)
        assert_locators_in_preview(path, result)

    @pytest.mark.parametrize("content", [b"PK\x03\x04bozuk", b""])
    def test_corrupt(self, tmp_path, content):
        path = tmp_path / "bozuk.pptx"
        path.write_bytes(content)
        assert ingest_files([path]).reports[0].status == "error"

    def test_zip_without_slides(self, tmp_path):
        path = tmp_path / "bos.pptx"
        path.write_bytes(generator().zip_bytes([("ppt/presentation.xml", b"<x/>")]))
        report = ingest_files([path]).reports[0]
        assert report.status == "error" and "slayt" in report.notes[0]


class TestOpenDocument:
    def test_odt_locators(self):
        path = FORMATS / "kart_tasarim.odt"
        result = ingest_files([path])
        used = source_locators(result, path.name)
        assert "§ Müşteri hattı · paragraf 2" in used  # liste maddesi paragraf sayılır
        assert any(u.startswith("tablo 1 (§ Boyut yüklemesi) · satır ") for u in used)
        assert_locators_in_preview(path, result)

    def test_ods_repeated_rows_are_bounded(self):
        started = time.monotonic()
        sheets = dict(opendocument.ods_sheets(FORMATS / "kart_sttm.ods"))
        assert time.monotonic() - started < 5
        assert len(sheets["STTM"]) == 6 and len(sheets["Bağımlılık"]) == 3  # 3 boş satır korunur, sondakiler yok
        assert sheets["STTM"][2:5] == [[], [], []]

    def test_ods_preview(self):
        path = FORMATS / "kart_sttm.ods"
        assert_locators_in_preview(path, ingest_files([path]))

    def test_encrypted_odf(self, tmp_path):
        manifest = (b'<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">'
                    b'<manifest:file-entry manifest:full-path="content.xml"><manifest:encryption-data/>'
                    b"</manifest:file-entry></manifest:manifest>")
        path = tmp_path / "sifreli.odt"
        path.write_bytes(generator().zip_bytes([("META-INF/manifest.xml", manifest), ("content.xml", b"\x93\x00")]))
        report = ingest_files([path]).reports[0]
        assert report.status == "error" and "şifre" in report.notes[0]

    @pytest.mark.parametrize("name", ["bozuk.odt", "bozuk.ods"])
    def test_corrupt(self, tmp_path, name):
        path = tmp_path / name
        path.write_bytes(generator().zip_bytes([("content.xml", b"<office:document-content")]))
        assert ingest_files([path]).reports[0].status == "error"


# --------------------------------------------------------------------------- #
# RTF ve HTML
# --------------------------------------------------------------------------- #


class TestRtf:
    def test_escapes_unicode_and_skipped_destinations(self):
        items = rtf_items(rb"{\rtf1\ansi\ansicpg1254{\fonttbl{\f0 Arial;}}{\*\generator Yazar;}"
                          rb"{\header BASLIK_A tablosu\par}"
                          rb"\uc1 T\u252?rk\'e7e \uc0\u351 k\u305\par"
                          rb"{\pict\bin6 }}}}}}\par}"
                          rb"Son\line sat\'fdr\tab x\par}")
        assert [i.text for i in items] == ["Türkçe şkı", "Son\nsatır x"]

    def test_heading_by_outline_level_and_numbering(self):
        items = rtf_items(rb"{\rtf1 \pard\outlinelevel0 Giri\'fe\par \pard 2.1 Y\'fckleme\par "
                          rb"\pard 1 STG_A tablosundan ODS_A tablosuna y\'fcklenir.\par}")
        assert [(i.kind, i.text) for i in items] == [
            ("heading", "Giriş"), ("heading", "2.1 Yükleme"),
            ("paragraph", "1 STG_A tablosundan ODS_A tablosuna yüklenir.")]

    def test_table_and_multi_paragraph_cell(self):
        items = rtf_items(rb"{\rtf1 \trowd\pard\intbl A\cell B\par devam\cell\row"
                          rb"\trowd\pard\intbl C\cell D\cell\row\pard sonra\par}")
        assert items[0].kind == "table" and items[0].rows == [["A", "B devam"], ["C", "D"]]
        assert items[1].text == "sonra"

    def test_unbalanced_braces_do_not_crash(self):
        assert [i.text for i in rtf_items(rb"{\rtf1 metin}}} ve {{ devam")] == ["metin ve devam"]

    def test_not_rtf(self, tmp_path):
        path = tmp_path / "x.rtf"
        path.write_text("düz metin", encoding="utf-8")
        report = ingest_files([path]).reports[0]
        assert report.status == "error" and "RTF" in report.notes[0]

    def test_preview_locators(self):
        path = FORMATS / "kart_notlari.rtf"
        assert_locators_in_preview(path, ingest_files([path]))


class TestHtml:
    def test_rowspan_colspan_alignment(self):
        items = html_items("<table><tr><th>A</th><th>B</th><th>C</th></tr>"
                           "<tr><td rowspan=2>x</td><td colspan=2>y</td></tr>"
                           "<tr><td>z</td><td>w</td></tr></table>")
        assert items[0].rows == [["A", "B", "C"], ["x", "y", ""], ["x", "z", "w"]]

    def test_layout_table_paragraphs_flow(self):
        items = html_items("<table><tr><td><h2>Başlık</h2><p>Bir</p><p>İki</p></td></tr></table><p>Üç</p>")
        assert [(i.kind, i.text) for i in items] == [
            ("heading", "Başlık"), ("paragraph", "Bir"), ("paragraph", "İki"), ("paragraph", "Üç")]

    def test_scripts_styles_head_ignored(self):
        items = html_items("<html><head><title>T</title><style>p{}</style><script>var a='STG_X';</script>"
                           "<body><p>Görünür</p><script>b()</script><noscript>gizli</noscript></body></html>")
        assert [i.text for i in items] == ["Görünür"]

    def test_unclosed_head_does_not_swallow_body(self):
        assert [i.text for i in html_items("<head><title>x<body><p>Gövde</p>")] == ["Gövde"]

    def test_charset_meta(self, tmp_path):
        path = tmp_path / "eski.htm"
        path.write_bytes('<meta charset="windows-1254"><p>STG_A tablosundan ODS_A tablosuna aktarılır.</p>'
                         .encode("cp1254"))
        assert summary(ingest_files([path]))["links"] == {("STG_A", "ODS_A")}

    def test_mhtml(self, tmp_path):
        path = tmp_path / "sayfa.mht"
        path.write_bytes(b"MIME-Version: 1.0\r\nContent-Type: multipart/related; boundary=\"s\"\r\n\r\n--s\r\n"
                         b"Content-Type: text/html; charset=utf-8\r\nContent-Transfer-Encoding: quoted-printable"
                         b"\r\n\r\n<p>STG_A tablosundan ODS_A tablosuna aktar=C4=B1l=C4=B1r.</p>\r\n--s--\r\n")
        assert summary(ingest_files([path]))["links"] == {("STG_A", "ODS_A")}

    def test_confluence_locators(self):
        path = FORMATS / "kart_confluence.html"
        result = ingest_files([path])
        used = source_locators(result, path.name)
        # 1: sayfa bilgisi (h1 altında), 2-3: ilk bölüm, 4: düzen tablosu içindeki paragraf
        assert "§ ODS ve olgu tablosu · paragraf 4" in used
        assert not {"HACK_A", "HACK_B"} & summary(result)["datasets"]
        assert_locators_in_preview(path, result)


# --------------------------------------------------------------------------- #
# E-posta
# --------------------------------------------------------------------------- #


def _eml(body: str, attachments: list[tuple[str, str, bytes]], subject: str = "Kart") -> bytes:
    return generator().eml_bytes({"From": "Ali Veli <ali.veli@banka.com>", "Subject": subject,
                                  "Date": "Tue, 15 Sep 2026 09:15:00 +0300"}, body, attachments)


class TestEmail:
    def test_attachments_are_separate_documents(self):
        result = ingest_files([FORMATS / "kart_degisiklik.eml"])
        inner = [r for r in result.reports if r.format != "arsiv"]
        assert {r.file for r in inner} == {"kart_degisiklik.govde.txt", "kart_sttm_ek.csv", "LOAD_MART_KART_OZET.sql"}
        assert {r.container for r in inner} == {"kart_degisiklik.eml"}
        assert "imza.png" not in {r.file for r in result.reports}  # satır içi görsel atlanır

    def test_body_file_has_headers(self, tmp_path):
        paths = _acma.expand_top(FORMATS / "kart_degisiklik.eml", tmp_path)
        body = next(p for p in paths if p.name.endswith(".govde.txt")).read_text(encoding="utf-8")
        head = body.split("Merhaba", 1)[0]
        assert "# E-posta: Kart mart yüklemesi değişikliği" in head
        assert "Kimden: Ayşe Yılmaz" in head and "@" not in head  # adres değil görünen ad
        assert "Tarih: Tue, 15 Sep 2026 09:15:00 +0300" in head

    def test_html_only_body(self, tmp_path):
        message = (b"From: a@b.c\r\nSubject: HTML\r\nMIME-Version: 1.0\r\n"
                   b"Content-Type: text/html; charset=utf-8\r\n\r\n"
                   b"<html><body><p>STG_A tablosundan ODS_A tablosuna aktar\xc4\xb1l\xc4\xb1r.</p></body></html>")
        path = tmp_path / "html.eml"
        path.write_bytes(message)
        result = ingest_files([path])
        body = only(result, "html.govde.html")
        assert body.format == "html" and body.container == "html.eml" and body.status == "ok"
        # 1: Konu, 2: Kimden (Tarih başlığı yok), 3: gövde paragrafı
        assert source_locators(result, "html.govde.html") == {"§ E-posta: HTML · paragraf 3"}

    def test_unsafe_attachment_name_stays_inside(self, tmp_path):
        path = tmp_path / "kotu.eml"
        path.write_bytes(_eml("x", [("../../../kacak.csv", "text/csv", b"Kaynak Tablo;Hedef Tablo\nA_1;B_1\n")]))
        work = tmp_path / "is"
        work.mkdir()
        paths = _acma.expand_top(path, work)
        assert all(p.resolve().is_relative_to(work.resolve()) for p in paths)
        assert not (tmp_path / "kacak.csv").exists()
        assert any(p.name.endswith("kacak.csv") for p in paths)

    def test_nested_zip_attachment_is_expanded(self, tmp_path):
        inner = generator().zip_bytes([("ic.sql", b"INSERT INTO ODS_A SELECT * FROM STG_A;")])
        path = tmp_path / "zipli.eml"
        path.write_bytes(_eml("Ekte.", [("paket.zip", "application/zip", inner)]))
        result = ingest_files([path])
        assert only(result, "ic.sql").container == "zipli.eml"
        assert summary(result)["links"] == {("STG_A", "ODS_A")}

    def test_unexpanded_eml_reads_body_and_notes_attachments(self, tmp_path):
        """Zorlanan okuyucu ya da derinlik sınırı: gövde okunur, ekler açılmaz ve raporlanır."""
        path = tmp_path / "posta.txt"
        path.write_bytes(_eml("STG_A tablosundan ODS_A tablosuna aktarılır.", [("a.csv", "text/csv", b"x;y\n")]))
        config = IngestConfig.from_settings({"okuyucu_zorla": {"posta.txt": "eposta"}})
        result = ingest_files([path], config=config)
        report = result.reports[0]
        assert report.format == "eposta" and report.status == "ok"
        assert any("1 ek" in n for n in report.notes)
        assert summary(result)["links"] == {("STG_A", "ODS_A")}

    def test_preview_lists_attachments(self):
        _, blocks = preview_document(FORMATS / "kart_degisiklik.eml")
        assert blocks[0]["locator"].startswith("E-posta: Kart mart")
        assert "kart_sttm_ek.csv" in blocks[-1]["text"]

    def test_msg_without_extra(self, tmp_path, scratch_registry, monkeypatch):
        monkeypatch.setitem(sys.modules, "extract_msg", None)
        REGISTRY.expanders.pop(".msg", None)
        path = tmp_path / "posta.msg"
        path.write_bytes(generator().CFB_SIGNATURE + b"\x00" * 512)
        report = ingest_files([path]).reports[0]
        assert report.status == "unsupported" and 'pip install "dwhflow[ofis]"' in " ".join(report.notes)

    def test_msg_with_extract_msg(self, tmp_path, scratch_registry, monkeypatch):
        closed = []

        class Attachment:
            def __init__(self, name, data):
                self.longFilename, self.data = name, data

        class Message:
            def __init__(self, subject, body, attachments):
                self.subject, self.sender, self.date = subject, "Ali Veli <ali@banka.com>", "2026-09-15 09:15"
                self.body, self.htmlBody, self.attachments = body, None, attachments

            def close(self):
                closed.append(self.subject)

        nested = Message("İletilen", "SRC_A tablosundan STG_A tablosuna aktarılır.", [])
        top = Message("Kart", "STG_A tablosundan ODS_A tablosuna aktarılır.",
                      [Attachment("ek.csv", b"Kaynak Tablo;Hedef Tablo\nODS_A;F_A\n"), Attachment("ic.msg", nested)])
        monkeypatch.setitem(sys.modules, "extract_msg", types.SimpleNamespace(openMsg=lambda path: top))
        register_expander((".msg",), _acma.expand_top)
        path = tmp_path / "posta.msg"
        path.write_bytes(b"sahte")
        result = ingest_files([path])
        inner = [r for r in result.reports if r.format != "arsiv"]
        assert {r.file for r in inner} == {"posta.govde.txt", "ek.csv", "ic.govde.txt"}
        assert {r.container for r in inner} == {"posta.msg"}
        assert summary(result)["links"] == {("STG_A", "ODS_A"), ("ODS_A", "F_A"), ("SRC_A", "STG_A")}
        assert closed == ["Kart"]


# --------------------------------------------------------------------------- #
# Zip
# --------------------------------------------------------------------------- #


def _zip(tmp_path: Path, entries: list[tuple[str, bytes]], name: str = "paket.zip", **kwargs) -> Path:
    path = tmp_path / name
    path.write_bytes(generator().zip_bytes(entries, **kwargs))
    return path


class TestZip:
    SQL = b"INSERT INTO ODS_A SELECT * FROM STG_A;"

    def test_inner_names_kept_and_duplicates_separated(self, tmp_path):
        path = _zip(tmp_path, [("a/yukle.sql", self.SQL), ("b/yukle.sql", self.SQL), ("c/YUKLE.sql", self.SQL)])
        work = tmp_path / "is"
        work.mkdir()
        paths = _acma.expand_top(path, work)
        assert sorted(p.name for p in paths) == ["YUKLE.sql", "yukle.sql", "yukle.sql"]
        assert len({p.resolve() for p in paths}) == 3
        assert all(p.resolve().is_relative_to(work.resolve()) for p in paths)

    def test_junk_skipped(self, tmp_path):
        path = _zip(tmp_path, [("x/", b""), ("__MACOSX/._a.sql", b"x"), (".git/config", b"x"), ("a/.DS_Store", b"x"),
                               ("Thumbs.db", b"x"), ("~$taslak.docx", b"x"), ("a.sql", self.SQL)])
        work = tmp_path / "is"
        work.mkdir()
        assert [p.name for p in _acma.expand_top(path, work)] == ["a.sql"]

    @pytest.mark.parametrize("bad", ["../kacak.sql", "a/../../kacak.sql", "/etc/kacak.sql", "C:/kacak.sql",
                                     "..\\kacak.sql"])
    def test_unsafe_paths_reject_archive(self, tmp_path, bad):
        path = _zip(tmp_path, [("iyi.sql", self.SQL), (bad, b"x")])
        result = ingest_files([path])
        assert [(r.file, r.status) for r in result.reports] == [("paket.zip", "error")]
        assert "yol" in result.reports[0].notes[0]
        assert not list(tmp_path.glob("kacak*")) and not (tmp_path.parent / "kacak.sql").exists()

    def test_file_count_cap(self, tmp_path, monkeypatch):
        monkeypatch.setattr(_acma, "MAX_FILES", 2)
        path = _zip(tmp_path, [(f"{i}.sql", self.SQL) for i in range(3)])
        report = ingest_files([path]).reports[0]
        assert report.status == "error" and "fazla dosya" in report.notes[0]

    def test_total_size_cap_counts_nested(self, tmp_path, monkeypatch):
        monkeypatch.setattr(_acma, "MAX_TOTAL_BYTES", 1000)
        inner = generator().zip_bytes([("b.txt", b"y" * 600)])
        path = _zip(tmp_path, [("a.txt", b"x" * 300), ("ic.zip", inner)])
        report = ingest_files([path]).reports[0]
        assert report.status == "error" and "MB" in report.notes[0]

    def test_ratio_guard(self, tmp_path):
        path = _zip(tmp_path, [("sifir.bin", b"\0" * (2 * 1024 * 1024))], compression=zipfile.ZIP_DEFLATED)
        report = ingest_files([path]).reports[0]
        assert report.status == "error" and "bombası" in report.notes[0]

    def test_turkish_dos_names(self, tmp_path):
        name = "Kart_İşlem.sql"
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            info = zipfile.ZipInfo("x" * len(name))
            archive.writestr(info, self.SQL)
        data = buffer.getvalue().replace(b"x" * len(name), name.encode("cp857"))  # UTF-8 bayraksız cp857 adı
        path = tmp_path / "dos.zip"
        path.write_bytes(data)
        work = tmp_path / "is"
        work.mkdir()
        assert [p.name for p in _acma.expand_top(path, work)] == [name]

    def test_corrupt_nested_archive_reported_by_its_reader(self, tmp_path):
        path = _zip(tmp_path, [("iyi.sql", self.SQL), ("bozuk.zip", b"PK\x03\x04 bozuk")])
        result = ingest_files([path])
        assert only(result, "iyi.sql").status == "ok"
        nested = only(result, "bozuk.zip")
        assert nested.status == "error" and nested.container == "paket.zip"

    def test_unsafe_nested_archive_not_extracted(self, tmp_path):
        inner = generator().zip_bytes([("../kacak.sql", self.SQL)])
        path = _zip(tmp_path, [("ic.zip", inner)])
        result = ingest_files([path])
        nested = only(result, "ic.zip")
        assert nested.status == "error" and "zip-slip" in nested.notes[0]

    def test_preview_lists_entries(self):
        key, blocks = preview_document(FORMATS / "kart_paket.zip")
        assert key == "arsiv" and "SP_ODS_KART.sql" in blocks[0]["text"] and "Thumbs.db" not in blocks[0]["text"]


# --------------------------------------------------------------------------- #
# Diyagramlar
# --------------------------------------------------------------------------- #


class TestDiagramRules:
    @pytest.mark.parametrize("name,hint,kind", [
        ("EXP_BDDK", "process", "process"),  # açık ipucu ad kuralını ezer
        ("LOAD_X", "dataset", "dataset"),
        ("STG_KART", None, "dataset"),  # ad kuralı katman tanıyor
        ("EXP_BDDK", None, "dataset"),  # EXP_ bir çıkış dosyası önekidir
        ("TPT_STG_KART", None, "process"),
        ("MFT_KART", None, "process"),
        ("Kart Yükleme İşi", None, "process"),  # "yükle" sözcüğü
        ("Müşteri Ana Verisi", None, "dataset"),
        ("KART_ISLEM.dat", None, "dataset"),
    ])
    def test_classify(self, name, hint, kind):
        assert _diyagram.classify(name, hint) == kind

    def test_classify_uses_configured_rules(self):
        config = IngestConfig.from_settings({"ad_kurallari": [{"tur": "onek", "deger": "JOB_", "katman": "ods"}]})
        with using(config):
            assert _diyagram.classify("JOB_TABLOSU", None) == "dataset"

    @pytest.mark.parametrize("label,expected", [
        ("LOAD_JOB", "LOAD_JOB"), ("SP_ODS_KART ile merge", "SP_ODS_KART"), ("günlük", None), ("", None),
        ("MFT", None), ("STG_A ve STG_B", None)])
    def test_label_process(self, label, expected):
        assert _diyagram.label_process(label) == expected

    @pytest.mark.parametrize("label,name", [
        ("Kaynak: SRC_KART_ISLEM", "SRC_KART_ISLEM"), ("STG_A\nstaging tablosu", "STG_A"),
        ("Müşteri Tablosu", "Müşteri Tablosu"), ("  ", None)])
    def test_node_name(self, label, name):
        found = _diyagram.node_name(label)
        assert (found[0] if found else None) == name

    def test_stop_labels_and_free_text(self):
        page = _diyagram.Page(prefix="sayfa 1")
        page.node("s", "Başla")
        page.node("a", "STG_A")
        page.node("b", "ODS_A")
        page.node("n", "Bu kutu yalnız bir açıklamadır")
        page.links = [_diyagram.Link("s", "a"), _diyagram.Link("a", "b", "SP_A")]
        from dwhflow.ingest.base import DocReport, ModelBuilder

        builder, report = ModelBuilder(), DocReport(file="d", reader="t")
        _diyagram.apply(builder, report, "d", [page], extractor="diyagram:test")
        model = builder.build()
        assert {d.name for d in model.datasets.values()} == {"STG_A", "ODS_A"}
        assert {p.name for p in model.processes.values()} == {"SP_A"}


class TestDrawio:
    def test_sniff_does_not_steal_other_xml(self, tmp_path):
        other = tmp_path / "zamanlayici.xml"
        other.write_text('<?xml version="1.0"?><DEFTABLE><FOLDER><JOB JOBNAME="A"/></FOLDER></DEFTABLE>',
                         encoding="utf-8")
        spec, _, _ = resolve(other, DEFAULT)
        assert spec is None or spec.key not in ("drawio", "bpmn", "sablon")

    def test_uncompressed_bare_model_and_placeholders(self, tmp_path):
        path = tmp_path / "ciplak.xml"
        path.write_text(
            '<mxGraphModel><root><mxCell id="0"/><mxCell id="1" parent="0"/>'
            '<UserObject id="a" label="%tablo%" placeholders="1" tablo="STG_A"><mxCell vertex="1" parent="1"/>'
            '</UserObject><mxCell id="b" value="ODS_A" vertex="1" parent="1"/>'
            '<mxCell id="e" value="SP_A" edge="1" parent="1" source="a" target="b"/></root></mxGraphModel>',
            encoding="utf-8")
        result = ingest_files([path])
        assert result.reports[0].format == "drawio"
        assert summary(result)["links"] == {("STG_A", "ODS_A")} and summary(result)["processes"] == {"SP_A"}

    def test_plain_label_keeps_angle_brackets(self, tmp_path):
        compressed = generator().drawio_compress(
            '<mxGraphModel><root><mxCell id="0"/><mxCell id="1" parent="0"/>'
            '<mxCell id="a" value="A&lt;B_TABLO" vertex="1" parent="1"/></root></mxGraphModel>')
        path = tmp_path / "x.drawio"
        path.write_text(f'<mxfile><diagram name="S">{compressed}</diagram></mxfile>', encoding="utf-8")
        _, blocks = preview_document(path)
        assert any(b["text"].startswith("A<B_TABLO") for b in blocks)

    @pytest.mark.parametrize("content", ["<mxfile><diagram>!!!bozuk!!!</diagram></mxfile>", "<mxfile><diagram>",
                                         "<baska/>"])
    def test_corrupt(self, tmp_path, content):
        path = tmp_path / "bozuk.drawio"
        path.write_text(content, encoding="utf-8")
        report = ingest_files([path]).reports[0]
        assert report.status == "error"

    @pytest.mark.parametrize("name", ["kart_akis.drawio", "kart_mimari_drawio.xml"])
    def test_preview_locators(self, name):
        path = FORMATS / name
        assert_locators_in_preview(path, ingest_files([path]))


class TestVisio:
    def test_reversed_connector_and_skipped_shapes(self):
        path = FORMATS / "kart_akis.vsdx"
        result = ingest_files([path])
        assert ("STG_KART_ISLEM", "ODS_KART_ISLEM") in summary(result)["links"]
        assert not {"GIZLI_ARKA_PLAN", "Kredi Kartı Veri Akışı"} & summary(result)["datasets"]
        assert "sayfa 2 · ok SRC_MUSTERI → STG_MUSTERI" in source_locators(result, path.name)  # arka plan sayılmaz
        assert_locators_in_preview(path, result)

    def test_not_a_visio_package(self, tmp_path):
        path = tmp_path / "x.vsdx"
        path.write_bytes(generator().zip_bytes([("a.txt", b"x")]))
        report = ingest_files([path]).reports[0]
        assert report.status == "error" and "Visio" in report.notes[0]


class TestBpmn:
    def test_sniffed_from_xml(self, tmp_path):
        path = tmp_path / "surec.xml"
        path.write_bytes((FORMATS / "kart_sureci.bpmn").read_bytes())
        assert resolve(path, DEFAULT)[0].key == "bpmn"

    def test_preview_locators(self):
        path = FORMATS / "kart_sureci.bpmn"
        result = ingest_files([path])
        assert "süreç 1 · ok TPT_STG_KART → SP_ODS_KART" in source_locators(result, path.name)  # geçit üzerinden
        assert_locators_in_preview(path, result)

    def test_without_process(self, tmp_path):
        path = tmp_path / "bos.bpmn"
        path.write_text('<definitions xmlns="http://www.omg.org/spec/BPMN/20100524/MODEL"/>', encoding="utf-8")
        assert ingest_files([path]).reports[0].status == "error"


class TestMermaid:
    def graph(self, text: str):
        page = mermaid.parse_mermaid(text).page
        return ({n.id: (n.label, n.hint) for n in page.nodes.values()},
                [(link.source, link.target, link.label) for link in page.links])

    def test_link_forms(self):
        nodes, links = self.graph("graph TD\nA-->B\nB --- C\nC -.-> D\nD ==> E\nE -- LOAD_X --> F\n"
                                  "F -->|SP_Y| G\nG -. SP_Z .-> H\nH ~~~ I\nI --x J\n")
        assert links == [("A", "B", ""), ("B", "C", ""), ("C", "D", ""), ("D", "E", ""), ("E", "F", "LOAD_X"),
                         ("F", "G", "SP_Y"), ("G", "H", "SP_Z"), ("I", "J", "")]

    def test_groups_chains_shapes_and_classes(self):
        nodes, links = self.graph('flowchart LR; A[(STG_A)] & B@{ shape: cyl, label: "STG_B" } --> P[[JOB_P]] --> '
                                  'C["ODS C"]:::veri\nclass P surec')
        assert links == [("A", "P", ""), ("B", "P", ""), ("P", "C", "")]
        assert nodes["A"] == ("STG_A", "dataset") and nodes["B"] == ("STG_B", "dataset")
        assert nodes["P"] == ("JOB_P", "process") and nodes["C"] == ("ODS C", "dataset")  # :::veri

    def test_unsupported_diagram_type(self, tmp_path):
        path = tmp_path / "er.mmd"
        path.write_text("erDiagram\n  MUSTERI ||--o{ KART : sahip", encoding="utf-8")
        report = ingest_files([path]).reports[0]
        assert report.status == "unsupported" and "flowchart" in report.notes[0]

    def test_unparsed_statements_noted(self, tmp_path):
        path = tmp_path / "x.mmd"
        path.write_text("graph LR\nSTG_A --> ODS_A\n)))((( \n", encoding="utf-8")
        result = ingest_files([path])
        assert summary(result)["links"] == {("STG_A", "ODS_A")}
        assert any("ayrıştırılamadı" in n for n in result.reports[0].notes)

    def test_sniff_txt(self, tmp_path):
        path = tmp_path / "akis.txt"
        path.write_text("%% not\nflowchart TD\n  STG_A --> ODS_A\n", encoding="utf-8")
        assert resolve(path, DEFAULT)[0].key == "mermaid"
        prose = tmp_path / "not.txt"
        prose.write_text("graph teorisi üzerine notlar\nSTG_A tablosundan ODS_A tablosuna yüklenir.", encoding="utf-8")
        assert resolve(prose, DEFAULT)[0].key == "metin"

    def test_preview_locators(self):
        path = FORMATS / "kart_akis.mmd"
        assert_locators_in_preview(path, ingest_files([path]))


class TestDot:
    def test_grammar(self):
        pages, undirected = dot_pages(
            '/* yorum */ strict digraph "g" { node [shape=cylinder]; A:p1 -> B:n [label=<<b>LOAD</b>_X>];\n'
            '# önişlemci satırı\n"DWH" + ".T" [label="DWH.T\\n(tablo)"]; B -> "DWH.T"; C [shape=component] }\n'
            "graph h { X -- Y }")
        first, second = pages
        assert undirected and first.prefix == "grafik 1" and second.prefix == "grafik 2"
        assert first.nodes["DWH.T"].label == "DWH.T\n(tablo)"
        assert [(link.source, link.target, link.label) for link in first.links] == [("A", "B", "LOAD_X"),
                                                                                    ("B", "DWH.T", "")]
        assert first.nodes["C"].hint == "process" and first.nodes["A"].hint == "dataset"

    def test_parse_error_is_reported(self, tmp_path):
        path = tmp_path / "bozuk.dot"
        path.write_text("digraph { A -> ", encoding="utf-8")
        report = ingest_files([path]).reports[0]
        assert report.status == "error" and "DOT" in report.notes[0]

    def test_word_template_with_dot_suffix(self, tmp_path, no_soffice):
        path = tmp_path / "sablon.dot"
        path.write_bytes(generator().doc_bytes([FLOW]))
        result = ingest_files([path])
        assert "Word" in result.reports[0].notes[0] and summary(result)["links"] == {("STG_A", "ODS_A")}

    def test_sniff_txt(self, tmp_path):
        path = tmp_path / "graf.txt"
        path.write_text("// açıklama\ndigraph G {\n STG_A -> ODS_A\n}\n", encoding="utf-8")
        assert resolve(path, DEFAULT)[0].key == "dot"

    def test_preview_locators(self):
        path = FORMATS / "kart_akis.gv"
        assert_locators_in_preview(path, ingest_files([path]))


# --------------------------------------------------------------------------- #
# Kurum şablonları
# --------------------------------------------------------------------------- #


@pytest.fixture(scope="module")
def template_config() -> IngestConfig:
    settings = yaml.safe_load((TEMPLATES / "sablonlar.yaml").read_text(encoding="utf-8"))
    return IngestConfig.from_settings(settings["okuma"])


class TestTemplates:
    def test_without_templates_unsupported(self):
        report = ingest_files([TEMPLATES / "kurum_etl.etl.xml"]).reports[0]
        assert report.status == "unsupported"

    def test_xml_template(self, template_config):
        path = TEMPLATES / "kurum_etl.etl.xml"
        result = ingest_files([path], config=template_config)
        report = result.reports[0]
        assert (report.format, report.status, report.detected) == ("sablon", "ok", "icerik")
        found = summary(result)
        assert found["links"] == {("SRC_KART_ISLEM", "KART_ISLEM_YYYYMMDD.dat"), ("STG_KART_ISLEM", "ODS_KART_ISLEM"),
                                  ("ODS_KART_ISLEM", "DWH.F_KART_ISLEM"), ("DWH.D_MUSTERI", "DWH.F_KART_ISLEM")}
        assert found["processes"] == {"EXT_KART", "SP_ODS_KART", "LOAD_F_KART_ISLEM"}
        ext = next(p for p in result.model.processes.values() if p.name == "EXT_KART")
        assert ext.description == "Kart işlemleri günlük extract"
        assert {(s.extractor, s.locator) for s in ext.sources} == {("sablon:Kurum ETL XML", "kayıt 1")}
        assert_locators_in_preview(path, result, template_config)

    def test_json_template(self, template_config):
        result = ingest_files([TEMPLATES / "kurum_isler.json"], config=template_config)
        found = summary(result)
        assert found["links"] == {("DWH.F_KART_ISLEM", "DWH.MART_KART_OZET"), ("DWH.D_MUSTERI", "DWH.MART_KART_OZET"),
                                  ("DWH.MART_KART_OZET", "EXP_BDDK_KART")}
        assert found["processes"] == {"LOAD_MART_KART_OZET", "EXP_BDDK"}

    def test_invalid_entries_noted_valid_ones_used(self, tmp_path):
        good = {"ad": "İyi", "dosya": "*.etl.xml", "tur": "xml", "kayit": ".//Mapping", "kaynak": "Source/@name",
                "hedef": "Target/@name", "surec": "@id"}
        config = IngestConfig.from_settings({"sablonlar": [
            "metin", {"ad": "Eksik", "dosya": "*.xml"}, {**good, "ad": "Tür", "tur": "csv"},
            {**good, "ad": "Yol", "kayit": ".//Mapping[@"}, {**good, "ad": "Süreç", "surec": 5}, good]})
        result = ingest_files([TEMPLATES / "kurum_etl.etl.xml"], config=config)
        notes = " ".join(result.reports[0].notes)
        for problem in ("okuma şablonu 1", "'Eksik' atlandı", "'Tür' atlandı", "'Yol' atlandı", "'Süreç' atlandı"):
            assert problem in notes
        assert result.reports[0].status == "ok" and len(summary(result)["links"]) == 4

    def test_sniff_is_strict(self, tmp_path, template_config):
        with using(template_config):
            assert not sablon.looks_like_template(b"<x/>", "kurum.etl.json")  # tür uymuyor
            assert not sablon.looks_like_template(b"<x/>", "baska.xml")  # desen uymuyor
            assert sablon.looks_like_template(b"<x/>", "Gece.ETL.XML")  # büyük/küçük harf duyarsız
            assert sablon.looks_like_template(b"<x/>", "gece.etl.xml")
        other = tmp_path / "kart_mimari_drawio.xml"
        other.write_bytes((FORMATS / "kart_mimari_drawio.xml").read_bytes())
        assert resolve(other, template_config)[0].key == "drawio"

    def test_template_can_claim_a_matching_name(self, tmp_path):
        """Kullanıcı açıkça desen verdiyse (ör. draw.io'ya benzeyen kendi XML'i) şablon devreye girebilir."""
        config = IngestConfig.from_settings({"sablonlar": [
            {"ad": "Kendi", "dosya": "ozel_*.xml", "tur": "xml", "kayit": "Kayit", "kaynak": "@a", "hedef": "@b"}]})
        path = tmp_path / "ozel_1.xml"
        path.write_text('<Kok><Kayit a="STG_A" b="ODS_A"/></Kok>', encoding="utf-8")
        assert summary(ingest_files([path], config=config))["links"] == {("STG_A", "ODS_A")}

    def test_forced_without_pattern_tries_by_type(self, tmp_path, template_config):
        path = tmp_path / "isler.dat"
        path.write_bytes((TEMPLATES / "kurum_isler.json").read_bytes())
        config = IngestConfig.from_settings({
            "sablonlar": list(template_config.templates), "okuyucu_zorla": {"isler.dat": "sablon"}})
        result = ingest_files([path], config=config)
        assert result.reports[0].detected == "zorla" and len(summary(result)["links"]) == 3
        assert any("uymadı" in n for n in result.reports[0].notes)

    def test_multiple_templates_prefix_locators(self, tmp_path):
        base = {"dosya": "*.json", "tur": "json", "kayit": "k[*]", "kaynak": "a", "hedef": "b"}
        config = IngestConfig.from_settings({"sablonlar": [{**base, "ad": "Bir"}, {**base, "ad": "İki"}]})
        path = tmp_path / "x.json"
        path.write_text(json.dumps({"k": [{"a": "STG_A", "b": "ODS_A"}]}), encoding="utf-8")
        result = ingest_files([path], config=config)
        assert source_locators(result, "x.json") == {"Bir · kayıt 1", "İki · kayıt 1"}

    @pytest.mark.parametrize("path,steps", [
        ("jobs[*].steps[*]", ["jobs", "*", "steps", "*"]), ("$.data.items[0]", ["data", "items", 0]), ("", [])])
    def test_json_steps(self, path, steps):
        assert sablon.json_steps(path) == steps

    def test_json_select_maps_keys_over_lists(self):
        document = {"sources": [{"name": "A"}, {"name": "B"}, "x"]}
        assert sablon.json_select(document, sablon.json_steps("sources.name")) == ["A", "B"]

    @pytest.mark.parametrize("content", ["<Kok><Kayit", "{bozuk"])
    def test_bad_content_is_error(self, tmp_path, content):
        suffix = ".xml" if content.startswith("<") else ".json"
        config = IngestConfig.from_settings({"sablonlar": [
            {"ad": "T", "dosya": f"*{suffix}", "tur": suffix[1:], "kayit": "Kayit", "kaynak": "@a", "hedef": "@b"}]})
        path = tmp_path / f"x{suffix}"
        path.write_text(content, encoding="utf-8")
        report = ingest_files([path], config=config).reports[0]
        assert report.status == "error"


# --------------------------------------------------------------------------- #
# Genel dönüştürücü (MarkItDown)
# --------------------------------------------------------------------------- #


class TestGenericFallback:
    def test_unknown_suffix_without_markitdown(self, tmp_path, monkeypatch):
        monkeypatch.setitem(sys.modules, "markitdown", None)
        path = tmp_path / "kitap.epub"
        path.write_bytes(b"PK\x03\x04 epub")
        report = ingest_files([path]).reports[0]
        assert report.status == "unsupported" and 'pip install "dwhflow[evrensel]"' in report.notes[0]

    def test_unknown_suffix_with_markitdown(self, tmp_path, monkeypatch):
        seen = []

        class MarkItDown:
            def convert_local(self, source):
                seen.append(source)
                return types.SimpleNamespace(markdown="# Başlık\n\nSTG_A tablosundan ODS_A tablosuna yüklenir.\n")

        monkeypatch.setitem(sys.modules, "markitdown", types.SimpleNamespace(MarkItDown=MarkItDown))
        path = tmp_path / "kitap.epub"
        path.write_bytes(b"PK\x03\x04 epub")
        result = ingest_files([path])
        report = result.reports[0]
        assert (report.format, report.detected, report.status) == ("genel", "yedek", "ok")
        assert seen == [str(path)] and summary(result)["links"] == {("STG_A", "ODS_A")}
        assert_locators_in_preview(path, result)

    def test_converter_failure_is_reported(self, tmp_path, monkeypatch):
        class MarkItDown:
            def convert_local(self, source):
                raise RuntimeError("desteklenmeyen dosya")

        monkeypatch.setitem(sys.modules, "markitdown", types.SimpleNamespace(MarkItDown=MarkItDown))
        path = tmp_path / "resim.png"
        path.write_bytes(b"\x89PNG\r\n\x1a\n")
        report = ingest_files([path]).reports[0]
        assert report.status == "unsupported" and "desteklenmeyen dosya" in " ".join(report.notes)


# --------------------------------------------------------------------------- #
# Örnek üreticisi
# --------------------------------------------------------------------------- #


def test_generator_reproduces_samples(tmp_path):
    """Üretici yeniden çalıştırılınca aynı bilgiyi taşıyan örnekler çıkar (altın beklentilerle)."""
    written = generator().generate(tmp_path / "bicim", tmp_path / "zor")
    for path in written:
        if path.parent.name != "bicim":
            continue
        expected = json.loads((FORMATS / f"{path.name}.beklenen.json").read_text(encoding="utf-8"))
        found = summary(ingest_files([path]))
        assert {d.upper() for d in found["datasets"]} == {d.upper() for d in expected["datasets"]}, path.name
        assert {p.upper() for p in found["processes"]} == {p.upper() for p in expected["processes"]}, path.name
