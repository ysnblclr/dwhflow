"""Dosya türüne göre okuyucular. Hepsi aynı ``ModelBuilder``'a yazar.

Her okuyucu ``(builder, report, path)`` alır ve LLM'e aday olabilecek, kuralla
çözülemeyen paragrafları ``[(konum, paragraf)]`` olarak döner.
"""

from __future__ import annotations

import csv
import io
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from ..model import ProcessKind, Schedule
from ..store import StoreError, load_model
from .base import DocReport, ModelBuilder, clean_name, ref
from .tables import interpret, rows_to_view
from .text import extract_paragraphs, read_text, split_markdown

Leftovers = list[tuple[str, str]]

MAX_ROWS_PER_SHEET = 20000
MAX_XML_BYTES = 60 * 1024 * 1024


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


# --------------------------------------------------------------------------- #
# Excel / CSV
# --------------------------------------------------------------------------- #


def read_excel(builder: ModelBuilder, report: DocReport, path: Path) -> Leftovers:
    import openpyxl

    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - bozuk dosya kullanıcıya raporlanır
        report.status = "error"
        report.note(f"Excel açılamadı: {exc}")
        return []
    leftovers: Leftovers = []
    try:
        for sheet in book.worksheets:
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_ROWS_PER_SHEET:
                    report.note(f"{sheet.title}: ilk {MAX_ROWS_PER_SHEET} satır okundu")
                    break
                rows.append(list(row))
            view = rows_to_view(rows, locator=sheet.title, row_locator="satır {n}")
            found = interpret(view, builder, report, file=path.name, reader="excel_sttm") if view else 0
            report.flows += found
            if not found:
                # Tablo olarak tanınmayan sayfada uzun metinler olabilir (doküman Excel'e yapıştırılmış).
                texts = [(f"{sheet.title} · satır {i + 1}", str(c)) for i, row in enumerate(rows)
                         for c in row if isinstance(c, str) and len(c) >= 30]
                result = extract_paragraphs(texts, builder, report, file=path.name)
                report.flows += result.flows
                leftovers.extend(result.leftovers)
    finally:
        book.close()
    return leftovers


def read_csv(builder: ModelBuilder, report: DocReport, path: Path) -> Leftovers:
    text = _decode(path.read_bytes())
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))[:MAX_ROWS_PER_SHEET]
    view = rows_to_view(rows, locator=path.stem, row_locator="satır {n}")
    if view:
        report.flows += interpret(view, builder, report, file=path.name, reader="excel_sttm")
    return []


# --------------------------------------------------------------------------- #
# Word (.docx) — ek kütüphane olmadan, doğrudan XML
# --------------------------------------------------------------------------- #

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _para_text(element: ET.Element) -> str:
    parts = []
    for node in element.iter():
        if node.tag == f"{_W}t" and node.text:
            parts.append(node.text)
        elif node.tag in (f"{_W}tab",):
            parts.append(" ")
        elif node.tag in (f"{_W}br", f"{_W}cr"):
            parts.append("\n")
    return "".join(parts).strip()


def _para_style(element: ET.Element) -> str:
    style = element.find(f"{_W}pPr/{_W}pStyle")
    return (style.get(f"{_W}val") or "") if style is not None else ""


def read_docx(builder: ModelBuilder, report: DocReport, path: Path) -> Leftovers:
    try:
        with zipfile.ZipFile(path) as archive:
            info = archive.getinfo("word/document.xml")
            if info.file_size > MAX_XML_BYTES:
                report.status = "error"
                report.note("Word dokümanı çok büyük")
                return []
            root = ET.fromstring(archive.read(info))
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
        report.status = "error"
        report.note(f"Word dosyası okunamadı (.docx olmalı): {exc}")
        return []

    body = root.find(f"{_W}body")
    if body is None:
        report.status = "empty"
        return []
    heading, table_no, paragraph_no = "", 0, 0
    paragraphs: list[tuple[str, str]] = []
    for element in body:
        if element.tag == f"{_W}p":
            text = _para_text(element)
            if not text:
                continue
            style = _para_style(element).lower()
            numbered = bool(re.match(r"^\d+(\.\d+)*\s", text)) and len(text) < 90
            if style.startswith(("heading", "baslik", "başlık", "title")) or numbered:
                heading = text[:60]
                continue
            paragraph_no += 1
            where = f"§ {heading} · paragraf {paragraph_no}" if heading else f"paragraf {paragraph_no}"
            paragraphs.append((where, text))
        elif element.tag == f"{_W}tbl":
            table_no += 1
            rows = []
            for row in element.iter(f"{_W}tr"):
                rows.append([" ".join(_para_text(p) for p in cell.iter(f"{_W}p")).strip()
                             for cell in row.iter(f"{_W}tc")])
            where = f"tablo {table_no}" + (f" (§ {heading})" if heading else "")
            view = rows_to_view(rows, locator=where, row_locator="satır {n}")
            if view:
                report.flows += interpret(view, builder, report, file=path.name, reader="docx_tablo")
    result = extract_paragraphs(paragraphs, builder, report, file=path.name)
    report.flows += result.flows
    return result.leftovers


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #


def read_pdf(builder: ModelBuilder, report: DocReport, path: Path) -> Leftovers:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            report.status = "error"
            report.note("PDF şifreli; okunamadı")
            return []
        pages = [(number, page.extract_text() or "") for number, page in enumerate(reader.pages, 1)]
    except Exception as exc:  # noqa: BLE001
        report.status = "error"
        report.note(f"PDF okunamadı: {exc}")
        return []
    if not any(text.strip() for _, text in pages):
        report.status = "empty"
        report.note("PDF'te metin yok (taranmış görüntü olabilir; OCR bu sürümde yok)")
        return []
    leftovers: Leftovers = []
    for number, text in pages:
        tables, paragraphs = split_markdown(text)
        for table in tables:
            table.locator = f"s. {number} · {table.locator}"
            report.flows += interpret(table, builder, report, file=path.name, reader="pdf_tablo")
        paragraphs = [(f"s. {number} · {where}", p) for where, p in paragraphs]
        # PDF satır sonları cümle ortasında kırılır; aynı sayfanın satırlarını birleştir.
        merged = [(f"s. {number}", " ".join(p for _, p in paragraphs))] if paragraphs else []
        result = extract_paragraphs(merged, builder, report, file=path.name)
        report.flows += result.flows
        leftovers.extend(result.leftovers)
    return leftovers


# --------------------------------------------------------------------------- #
# Düz metin / Markdown
# --------------------------------------------------------------------------- #


def read_textfile(builder: ModelBuilder, report: DocReport, path: Path) -> Leftovers:
    result = read_text(_decode(path.read_bytes()), builder, report, file=path.name)
    report.flows += result.flows
    return result.leftovers


# --------------------------------------------------------------------------- #
# SQL
# --------------------------------------------------------------------------- #

_DOT_COMMAND = re.compile(r"^\s*\.[A-Za-z]+.*$", re.MULTILINE)  # BTEQ .LOGON, .IF …
_DIALECTS = (None, "teradata", "oracle", "tsql", "postgres", "snowflake")


def _table_name(table) -> str:
    parts = [p for p in (table.catalog, table.db, table.name) if p]
    return ".".join(parts)


def _statements(text: str) -> list[str]:
    """Noktalı virgülle böler; tırnak içindeki ; atlanır."""
    out, buffer, quote = [], [], None
    for char in text:
        if quote:
            if char == quote:
                quote = None
        elif char in ("'", '"'):
            quote = char
        elif char == ";":
            out.append("".join(buffer))
            buffer = []
            continue
        buffer.append(char)
    out.append("".join(buffer))
    return [s.strip() for s in out if s.strip()]


def _lineage(statement: str) -> tuple[str | None, list[str], str] | None:
    """(hedef, kaynaklar, ifade türü) ya da None."""
    import sqlglot
    from sqlglot import exp

    for dialect in _DIALECTS:
        try:
            tree = sqlglot.parse_one(statement, read=dialect)
        except Exception:  # noqa: BLE001 - sonraki lehçeyi dene
            continue
        if tree is None:
            continue
        ctes = {cte.alias_or_name.upper() for cte in tree.find_all(exp.CTE)}
        target_node = None
        kind = type(tree).__name__.lower()
        if isinstance(tree, exp.Insert):
            target_node = tree.this.find(exp.Table) if not isinstance(tree.this, exp.Table) else tree.this
        elif isinstance(tree, exp.Create) and tree.expression is not None:
            target_node = tree.this.find(exp.Table) if not isinstance(tree.this, exp.Table) else tree.this
        elif isinstance(tree, (exp.Merge, exp.Update)):
            target_node = tree.this if isinstance(tree.this, exp.Table) else tree.this.find(exp.Table)
        else:
            return None
        if target_node is None:
            return None
        target = _table_name(target_node)
        sources = []
        for table in tree.find_all(exp.Table):
            name = _table_name(table)
            if table is target_node or not name or name.upper() in ctes or name.upper() == target.upper():
                continue
            if name not in sources:
                sources.append(name)
        return target, sources, kind
    return None


def read_sql(builder: ModelBuilder, report: DocReport, path: Path) -> Leftovers:
    text = _DOT_COMMAND.sub("", _decode(path.read_bytes()))
    statements = _statements(text)
    found = []
    unparsed = 0
    for index, statement in enumerate(statements, start=1):
        lineage = _lineage(statement)
        if lineage is None:
            if re.match(r"^\s*(insert|create|merge|update)\b", statement, re.IGNORECASE):
                unparsed += 1
            continue
        found.append((index, *lineage))
    if unparsed:
        report.note(f"{unparsed} veri yazan ifade ayrıştırılamadı (dinamik SQL ya da lehçe farkı)")
    if not found:
        report.status = "empty" if not unparsed else "error"
        report.note("veri akışı içeren SQL ifadesi bulunamadı")
        return []
    many = len(found) > 1
    for index, target, sources, kind in found:
        locator = f"ifade {index} ({kind})"
        source = ref(path.name, locator, "sql:sqlglot", confidence=1.0)
        process_name = f"{path.stem} #{index}" if many else path.stem
        process = builder.process(process_name, source, kind=ProcessKind.SQL_SCRIPT, tool="SQL")
        target_key = builder.dataset(target, source)
        builder.writes(process, target_key, source)
        for name in sources:
            builder.reads(builder.dataset(name, source), process, source)
            report.flows += 1
    return []


# --------------------------------------------------------------------------- #
# AutoSys JIL
# --------------------------------------------------------------------------- #

_JIL_ATTR = re.compile(r"(\w+)\s*:\s*(.*?)(?=\s+\w+\s*:\s|$)")
_JIL_COND = re.compile(r"\b(?:s|success|d|done)\s*\(\s*([\w.#-]+)", re.IGNORECASE)


def read_jil(builder: ModelBuilder, report: DocReport, path: Path) -> Leftovers:
    text = _decode(path.read_bytes())
    jobs: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    for number, raw in enumerate(text.splitlines(), start=1):
        line = raw.split("/*")[0].strip()
        if not line:
            continue
        for key, value in _JIL_ATTR.findall(line):
            key = key.lower()
            if key == "insert_job":
                current = {"insert_job": value.strip(), "_line": str(number)}
                jobs.append(current)
            elif current is not None:
                current[key] = value.strip().strip('"')
    if not jobs:
        report.status = "empty"
        report.note("insert_job tanımı bulunamadı")
        return []
    for job in jobs:
        source = ref(path.name, f"satır {job['_line']}", "scheduler:autosys", confidence=1.0)
        times = [t.strip() for t in job.get("start_times", "").split(",") if t.strip()]
        key = builder.process(
            job["insert_job"], source, kind=ProcessKind.BATCH_JOB, tool="AutoSys",
            description=job.get("description"),
            schedule=Schedule(start_times=times, days_of_week=[
                d.strip() for d in job.get("days_of_week", "").split(",") if d.strip()
            ]) if times else None,
        )
        report.processes += 1
        for previous in _JIL_COND.findall(job.get("condition", "")):
            builder.precedes(builder.process(previous, source, kind=ProcessKind.BATCH_JOB), key, source)
            report.flows += 1
    return []


# --------------------------------------------------------------------------- #
# YAML model
# --------------------------------------------------------------------------- #


def read_yaml(builder: ModelBuilder, report: DocReport, path: Path) -> Leftovers:
    try:
        model = load_model(path)
    except StoreError as exc:
        report.status = "error"
        report.note(f"DWHFlow model dosyası değil ya da geçersiz: {str(exc).splitlines()[0][:200]}")
        return []
    builder.merge_model(model)
    report.datasets, report.processes, report.flows = (
        len(model.datasets), len(model.processes), len(model.edges))
    return []


__all__ = [
    "clean_name",
    "read_csv",
    "read_docx",
    "read_excel",
    "read_jil",
    "read_pdf",
    "read_sql",
    "read_textfile",
    "read_yaml",
]
