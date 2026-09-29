"""Doküman önizlemesi: belgeyi okuyucuların kullandığı **aynı konum etiketleriyle** bloklara böler.

Doküman inceleme ekranı, modeldeki bir öğenin kaynağındaki ``locator`` ve ``evidence``
değerleriyle bu blokları eşleştirip kanıtı vurgular. Bu yüzden konum etiketleri
okuyuculardakiyle birebir aynı üretilir (ör. Word: ``§ Başlık · paragraf 3``,
Excel: sayfa adı + ``satır n``, PDF: ``s. 4``).

Blok biçimi::

    {"locator": "§ Yükleme · paragraf 3", "kind": "paragraph", "text": "..."}
    {"locator": "STTM", "kind": "table", "headers": [...], "rows": [[...], ...],
     "row_locator": "satır {n}"}
    {"locator": "satır 1", "kind": "code", "text": "...", "line_start": 1}
"""

from __future__ import annotations

import csv
import io
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from .readers import _W, MAX_ROWS_PER_SHEET, MAX_XML_BYTES, _decode, _para_style, _para_text, _statements
from .registry import Block
from .tables import rows_to_view
from .text import split_markdown

MAX_BLOCKS = 5000
MAX_PREVIEW_ROWS = 2000
MAX_TEXT_BYTES = 2 * 1024 * 1024


def _table_block(rows: list[list[object]], locator: str) -> Block | None:
    view = rows_to_view([list(r) for r in rows], locator=locator)
    if view is None:
        return None
    return {"locator": locator, "kind": "table", "headers": view.headers,
            "rows": view.rows[:MAX_PREVIEW_ROWS], "row_locator": view.row_locator,
            "truncated": len(view.rows) > MAX_PREVIEW_ROWS}


def preview_excel(path: Path) -> list[Block]:
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    blocks: list[Block] = []
    try:
        for sheet in book.worksheets:
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_ROWS_PER_SHEET:
                    break
                rows.append(["" if c is None else str(c) for c in row])
            block = _table_block(rows, sheet.title)
            blocks.append(block or {"locator": sheet.title, "kind": "note",
                                    "text": "Sayfada tablo bulunamadı."})
    finally:
        book.close()
    return blocks


def preview_csv(path: Path) -> list[Block]:
    text = _decode(path.read_bytes()[:MAX_TEXT_BYTES])
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ";" if text[:4096].count(";") > text[:4096].count(",") else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))[:MAX_ROWS_PER_SHEET]
    block = _table_block(rows, path.stem)
    return [block] if block else [{"locator": path.stem, "kind": "note", "text": "Tablo bulunamadı."}]


def preview_docx(path: Path) -> list[Block]:
    with zipfile.ZipFile(path) as archive:
        info = archive.getinfo("word/document.xml")
        if info.file_size > MAX_XML_BYTES:
            return [{"locator": "", "kind": "note", "text": "Doküman önizleme için çok büyük."}]
        root = ET.fromstring(archive.read(info))
    body = root.find(f"{_W}body")
    blocks: list[Block] = []
    if body is None:
        return blocks
    heading, table_no, paragraph_no = "", 0, 0
    for element in body:
        if element.tag == f"{_W}p":
            text = _para_text(element)
            if not text:
                continue
            style = _para_style(element).lower()
            numbered = bool(re.match(r"^\d+(\.\d+)*\s", text)) and len(text) < 90
            if style.startswith(("heading", "baslik", "başlık", "title")) or numbered:
                heading = text[:60]
                blocks.append({"locator": f"§ {heading}", "kind": "heading", "text": text})
                continue
            paragraph_no += 1
            where = f"§ {heading} · paragraf {paragraph_no}" if heading else f"paragraf {paragraph_no}"
            blocks.append({"locator": where, "kind": "paragraph", "text": text})
        elif element.tag == f"{_W}tbl":
            table_no += 1
            rows = [[" ".join(_para_text(p) for p in cell.iter(f"{_W}p")).strip()
                     for cell in row.iter(f"{_W}tc")] for row in element.iter(f"{_W}tr")]
            where = f"tablo {table_no}" + (f" (§ {heading})" if heading else "")
            block = _table_block(rows, where)
            if block:
                blocks.append(block)
        if len(blocks) >= MAX_BLOCKS:
            break
    return blocks


def preview_pdf(path: Path) -> list[Block]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        return [{"locator": "", "kind": "note", "text": "PDF şifreli."}]
    blocks: list[Block] = []
    for number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            blocks.append({"locator": f"s. {number}", "kind": "paragraph", "text": text})
        if len(blocks) >= MAX_BLOCKS:
            break
    return blocks


def preview_markdown(path: Path) -> list[Block]:
    text = _decode(path.read_bytes()[:MAX_TEXT_BYTES])
    tables, paragraphs = split_markdown(text)
    blocks: list[Block] = [{"locator": where, "kind": "paragraph", "text": body}
                           for where, body in paragraphs]
    for table in tables:
        blocks.append({"locator": table.locator, "kind": "table", "headers": table.headers,
                       "rows": table.rows[:MAX_PREVIEW_ROWS], "row_locator": table.row_locator})
    return blocks[:MAX_BLOCKS]


def preview_sql(path: Path) -> list[Block]:
    text = _decode(path.read_bytes()[:MAX_TEXT_BYTES])
    text = re.sub(r"^\s*\.[A-Za-z]+.*$", "", text, flags=re.MULTILINE)
    return [{"locator": f"ifade {index}", "kind": "code", "text": statement}
            for index, statement in enumerate(_statements(text), start=1)][:MAX_BLOCKS]


def preview_code(path: Path) -> list[Block]:
    """Satır numaralı düz metin (JIL, YAML, betikler): konumlar ``satır n`` biçimindedir."""
    text = _decode(path.read_bytes()[:MAX_TEXT_BYTES])
    return [{"locator": "satır 1", "kind": "code", "text": text, "line_start": 1}]


def preview_any(path: Path) -> list[Block]:
    """Önizleyicisi olmayan biçimler: metin gibi görünüyorsa göster, değilse not düş."""
    data = path.read_bytes()[:MAX_TEXT_BYTES]
    if b"\x00" in data[:4096]:
        return [{"locator": "", "kind": "note",
                 "text": "Bu biçim için önizleme yok (ikili dosya)."}]
    return [{"locator": "satır 1", "kind": "code", "text": _decode(data), "line_start": 1}]
